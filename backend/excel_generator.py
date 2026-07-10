import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

def generate_excel_file(transactions: list, column_mapping: dict, column_sequence: list) -> bytes:
    """
    Transforms parsed transaction list into a highly styled Excel file based on 
    user custom column mappings and layout sequence.
    
    Args:
        transactions: List of dicts representing transactions.
        column_mapping: Dict mapping standard internal keys ('date', 'description', 'debit', 'credit', 'amount')
                        to user-custom headers (e.g., {'date': 'Tx Date', 'description': 'Vendor Name', ...})
        column_sequence: List of custom headers indicating the order of columns.
        
    Returns:
        bytes: The binary content of the generated Excel file.
    """
    # 1. Flatten transactions into standard format
    rows = []
    for tx in transactions:
        tx_type = tx.get("type", "debit").lower()
        amount = float(tx.get("amount", 0.0))
        
        row = {
            "date": tx.get("date", ""),
            "description": tx.get("description", ""),
            "amount": amount if tx_type == "credit" else -amount,
            "debit": amount if tx_type == "debit" else "",
            "credit": amount if tx_type == "credit" else ""
        }
        rows.append(row)
        
    df_raw = pd.DataFrame(rows)
    
    # 2. Map standard columns to custom columns
    df_mapped = pd.DataFrame()
    for std_key, custom_name in column_mapping.items():
        if std_key in df_raw.columns and custom_name:
            df_mapped[custom_name] = df_raw[std_key]
            
    # 3. Ensure all columns in sequence exist (fill with empty if missing)
    for col in column_sequence:
        if col not in df_mapped.columns:
            df_mapped[col] = ""
            
    # 4. Reorder according to the user sequence
    if column_sequence:
        df_mapped = df_mapped[column_sequence]
        
    # 5. Write to openpyxl workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.views.sheetView[0].showGridLines = True
    
    # Write headers
    headers = list(df_mapped.columns)
    ws.append(headers)
    
    # Write data
    for r in df_mapped.itertuples(index=False):
        ws.append(list(r))
        
    # 6. Apply Premium Formatting and Styles
    # Color palette: Elegant Slate Dark (#2D3748) for headers, white text
    header_fill = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # Style header row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    ws.row_dimensions[1].height = 28
    
    # Style data rows
    # Identify column indexes that represent numerical financial values
    financial_cols = []
    # Reverse lookup map
    reverse_map = {v: k for k, v in column_mapping.items()}
    
    for idx, header in enumerate(headers):
        std_key = reverse_map.get(header)
        if std_key in ["amount", "debit", "credit"]:
            financial_cols.append(idx + 1)
            
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            
            # Format numbers for debit/credit/amount
            if col_idx in financial_cols:
                if cell.value != "" and cell.value is not None:
                    try:
                        cell.value = float(cell.value)
                        cell.number_format = "$#,##0.00"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    except ValueError:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            elif headers[col_idx-1].lower() in ["date", "tx date", "transaction date"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        
        # Determine maximum contents length
        for cell in col:
            val = str(cell.value or '')
            # If numerical with currency formatting, give it padding
            if cell.column in financial_cols and isinstance(cell.value, (int, float)):
                val = f"${cell.value:,.2f}"
            max_len = max(max_len, len(val))
            
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Save to binary output stream
    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    return excel_stream.getvalue()
